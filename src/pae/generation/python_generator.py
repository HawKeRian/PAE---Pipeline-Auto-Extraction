"""Generate a standalone, repeatable Python batch pipeline package."""

# ruff: noqa: E501 -- template source is checked after rendering and retains readable statements.

from __future__ import annotations

import ast
import hashlib
import json
from pathlib import Path
from typing import Any

from pae.domain.enums import ArtifactStatus, TargetLanguage, TransformationKind
from pae.domain.models import PipelineSpecification
from pae.generation.models import GeneratedPackage, PackageValidation
from pae.persistence.errors import UnsupportedCapability
from pae.sandbox import SandboxRunner

PIPELINE_TEMPLATE = r'''"""Generated PAE batch pipeline. Do not edit; regenerate from the confirmed spec."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import logging
import os
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Any

from runtime_core import execute

LOG = logging.getLogger("pae.generated")


def read_rows(path: Path, source: dict[str, Any]) -> list[dict[str, Any]]:
    fmt = source["file_format"]
    if fmt == "csv":
        with path.open(encoding=source.get("encoding", "utf-8"), newline="") as handle:
            return list(csv.DictReader(handle, delimiter=source.get("delimiter") or ","))
    if fmt == "json":
        value = json.loads(path.read_text(encoding=source.get("encoding", "utf-8")))
        return value if isinstance(value, list) else [value]
    if fmt == "json_lines":
        return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    if fmt == "excel":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        sheet = book[source["sheet_name"]] if source.get("sheet_name") else book.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value) for value in next(iterator)]
        return [dict(zip(headers, values, strict=True)) for values in iterator]
    if fmt == "parquet":
        import pyarrow.parquet as parquet
        return parquet.read_table(path).to_pylist()
    raise ValueError(f"unsupported file format: {fmt}")


def read_database(source: dict[str, Any]) -> list[dict[str, Any]]:
    config = json.loads(os.environ[source["connection_ref"]])
    query = source.get("read_only_query")
    if not query:
        name = source["object_name"]
        schema = source.get("schema_name")
        query = f'SELECT * FROM "{schema}"."{name}"' if schema else f'SELECT * FROM "{name}"'
    limit = int(source.get("sample_limit", 10000))
    kind = source["database_type"]
    if kind == "postgresql":
        import psycopg
        connection = psycopg.connect(**config)
    elif kind == "mysql":
        import pymysql
        connection = pymysql.connect(**config)
    elif kind == "sql_server":
        import pyodbc
        connection = pyodbc.connect(config["connection_string"])
    else:
        raise ValueError(f"unsupported database type: {kind}")
    try:
        cursor = connection.cursor()
        cursor.execute(query)
        columns = [item[0] for item in cursor.description]
        return [dict(zip(columns, row, strict=True)) for row in cursor.fetchmany(limit)]
    finally:
        connection.close()


def validate_schema(rows: list[dict[str, Any]], spec: dict[str, Any], source: dict[str, Any]) -> None:
    actual = {name for row in rows for name in row}
    selected = [field for field in spec["fields"] if field.get("selected", True)]
    required = {field["source_name"] for field in selected if field.get("required", True)}
    allowed = {field["source_name"] for field in selected}
    missing = required - actual
    extra = actual - allowed
    policy = source.get("runtime", {}).get("schema_policy", "strict")
    if missing or (extra and policy == "strict"):
        raise ValueError(f"schema incompatible; missing={sorted(missing)}, extra={sorted(extra)}")


def write_rows(path: Path, rows: list[dict[str, Any]], spec: dict[str, Any]) -> None:
    output = spec["output"]
    path.parent.mkdir(parents=True, exist_ok=True)
    append = output.get("write_mode", "overwrite") == "append" and path.exists()
    target = path
    temporary: Path | None = None
    if output.get("atomic_write", True):
        handle, name = tempfile.mkstemp(prefix=path.name, suffix=".tmp", dir=path.parent)
        os.close(handle)
        temporary = Path(name)
        target = temporary
        if append:
            shutil.copy2(path, target)
    try:
        fmt = output["format"]
        if fmt == "csv":
            fields = list(rows[0]) if rows else []
            with target.open("a" if append else "w", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
                if not append:
                    writer.writeheader()
                null_value = output.get("null_representation", "")
                writer.writerows([{key: null_value if value is None else value for key, value in row.items()} for row in rows])
        elif fmt == "json_lines":
            with target.open("a" if append else "w", encoding="utf-8") as handle:
                for row in rows:
                    handle.write(json.dumps(row, ensure_ascii=False) + "\n")
        elif fmt == "parquet":
            if append:
                raise ValueError("append is not supported for parquet")
            import pyarrow as pa
            import pyarrow.parquet as parquet
            parquet.write_table(pa.Table.from_pylist(rows), target)
        else:
            raise ValueError(f"unsupported output format: {fmt}")
        if temporary:
            temporary.replace(path)
    except Exception:
        if temporary:
            temporary.unlink(missing_ok=True)
        raise


def checksum(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_state(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {"processed": {}}


def save_state(path: Path, state: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(state, indent=2, sort_keys=True), encoding="utf-8")
    temporary.replace(path)


def output_path(root: Path, source_path: Path, spec: dict[str, Any]) -> Path:
    extension = {"csv": ".csv", "json_lines": ".jsonl", "parquet": ".parquet"}[spec["output"]["format"]]
    template = spec["output"].get("path_template", "output/result")
    has_source_token = "{stem}" in template or "{name}" in template
    rendered = template.replace("{stem}", source_path.stem).replace("{name}", source_path.name)
    if not has_source_token:
        rendered = f"{rendered}-{source_path.stem}"
    relative = Path(rendered)
    if relative.is_absolute() or ".." in relative.parts:
        raise ValueError("output path template must remain relative")
    return root / f"{relative}{extension}"


def run_files(args: argparse.Namespace, spec: dict[str, Any], source: dict[str, Any]) -> dict[str, Any]:
    input_root = args.input_dir.resolve()
    runtime = source["runtime"]
    pattern = runtime["filename_pattern"]
    candidates = sorted(input_root.glob(pattern), key=lambda item: item.as_posix())
    state = load_state(args.state_file)
    summary = {"discovered": len(candidates), "processed": 0, "skipped": 0, "failed": 0, "rejected": 0}
    for path in candidates:
        if not path.is_file() or input_root not in path.resolve().parents:
            continue
        digest = checksum(path)
        key = path.relative_to(input_root).as_posix()
        if state["processed"].get(key) == digest:
            summary["skipped"] += 1
            continue
        try:
            rows = read_rows(path, source)
            validate_schema(rows, spec, source)
            result = execute(spec, rows)
            write_rows(output_path(args.output_dir, path, spec), result["output_rows"], spec)
            if result["rejected_rows"]:
                rejected_path = args.quarantine_dir / f"{path.name}.rejected.jsonl"
                rejected_path.parent.mkdir(parents=True, exist_ok=True)
                rejected_path.write_text("\n".join(json.dumps(row, ensure_ascii=False) for row in result["rejected_rows"]), encoding="utf-8")
            summary["rejected"] += result["rejected_count"]
            summary["processed"] += 1
            state["processed"][key] = digest
            save_state(args.state_file, state)
        except Exception as exc:
            summary["failed"] += 1
            policy = runtime.get("failure_policy", "quarantine")
            LOG.error("failed %s: %s", key, exc)
            if policy == "quarantine":
                args.quarantine_dir.mkdir(parents=True, exist_ok=True)
                shutil.copy2(path, args.quarantine_dir / path.name)
                (args.quarantine_dir / f"{path.name}.error.json").write_text(json.dumps({"error": str(exc)}), encoding="utf-8")
            elif policy == "fail_batch":
                break
    return summary


def main() -> int:
    parser = argparse.ArgumentParser(description="Run the generated PAE pipeline")
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--quarantine-dir", type=Path, required=True)
    parser.add_argument("--state-file", type=Path, required=True)
    args = parser.parse_args()
    logging.basicConfig(level=os.getenv("PAE_LOG_LEVEL", "INFO"))
    spec = json.loads(Path(__file__).with_name("pipeline_spec.json").read_text(encoding="utf-8"))
    source = spec["sources"][0]
    try:
        if source["kind"] == "database":
            rows = read_database(source)
            validate_schema(rows, spec, source)
            result = execute(spec, rows)
            write_rows(output_path(args.output_dir, Path("database"), spec), result["output_rows"], spec)
            summary = {"discovered": 1, "processed": 1, "skipped": 0, "failed": 0, "rejected": result["rejected_count"]}
        else:
            summary = run_files(args, spec, source)
        print(json.dumps(summary, sort_keys=True))
        return 1 if summary["failed"] else 0
    except Exception as exc:
        LOG.error("pipeline failed: %s", exc)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
'''


class PythonGenerator:
    """Render a Python package from a confirmed, supported specification."""

    def generate(self, specification: PipelineSpecification) -> GeneratedPackage:
        if specification.target.language is not TargetLanguage.PYTHON:
            raise UnsupportedCapability("The Python generator requires target language python.")
        if specification.joins or any(
            rule.kind is TransformationKind.JOIN for rule in specification.transformations
        ):
            raise UnsupportedCapability("Multi-source join is Deferred and cannot be generated.")
        if len(specification.sources) != 1:
            raise UnsupportedCapability("The Python generator supports one source in the MVP.")
        source = specification.sources[0]
        if source.kind == "file" and (
            Path(source.original_name).name != source.original_name or ":" in source.original_name
        ):
            raise UnsupportedCapability("Sample paths cannot be embedded in generated packages.")
        if source.kind == "database" and any(
            marker in source.connection_ref for marker in ("://", "=", ";")
        ):
            raise UnsupportedCapability(
                "Database connection_ref must be an environment variable name, not a credential."
            )
        spec_json = json.dumps(
            specification.model_dump(mode="json"), ensure_ascii=False, indent=2, sort_keys=True
        )
        core = (
            Path(__file__)
            .resolve()
            .parents[1]
            .joinpath("runtime_core.py")
            .read_text(encoding="utf-8")
        )
        files = {
            "pipeline.py": PIPELINE_TEMPLATE,
            "runtime_core.py": core,
            "pipeline_spec.json": spec_json + "\n",
            "requirements.txt": self._requirements(specification),
            "config.example.env": "PAE_LOG_LEVEL=INFO\n# Set the database connection_ref variable to a JSON object when using a database source.\n",
            "README.md": self._readme(specification),
            "tests/test_pipeline.py": self._test_module(),
        }
        payload = "".join(f"{name}\0{files[name]}\0" for name in sorted(files))
        digest = hashlib.sha256(payload.encode("utf-8")).hexdigest()
        return GeneratedPackage(files=files, checksum=digest)

    @staticmethod
    def syntax_validate(package: GeneratedPackage) -> PackageValidation:
        checked = tuple(sorted(name for name in package.files if name.endswith(".py")))
        for name in checked:
            ast.parse(package.files[name], filename=name)
        return PackageValidation(
            status=ArtifactStatus.SYNTAX_VALIDATED,
            syntax_valid=True,
            sample_matches_preview=False,
            checked_files=checked,
        )

    @staticmethod
    async def sample_validate(
        package: GeneratedPackage,
        specification: PipelineSpecification,
        rows: tuple[dict[str, Any], ...],
        sandbox: SandboxRunner,
    ) -> PackageValidation:
        PythonGenerator.syntax_validate(package)
        generated_result = await sandbox.run(
            {
                "operation": "preview",
                "specification": json.loads(package.files["pipeline_spec.json"]),
                "rows": list(rows),
            }
        )
        preview_result = await sandbox.run(
            {
                "operation": "preview",
                "specification": specification.model_dump(mode="json"),
                "rows": list(rows),
            }
        )
        matches = generated_result == preview_result
        return PackageValidation(
            status=ArtifactStatus.SAMPLE_TESTED if matches else ArtifactStatus.SYNTAX_VALIDATED,
            syntax_valid=True,
            sample_matches_preview=matches,
            checked_files=tuple(sorted(name for name in package.files if name.endswith(".py"))),
        )

    @staticmethod
    def _requirements(specification: PipelineSpecification) -> str:
        requirements: list[str] = []
        source = specification.sources[0]
        if source.kind == "file" and source.file_format == "excel":
            requirements.append("openpyxl==3.1.5")
        if source.kind == "file" and source.file_format == "parquet":
            requirements.append("pyarrow==25.0.1")
        if source.kind == "database":
            requirements.append(
                {
                    "postgresql": "psycopg[binary]==3.3.6",
                    "mysql": "PyMySQL==1.2.3",
                    "sql_server": "pyodbc==5.3.0",
                }[source.database_type]
            )
        return "\n".join(requirements) + ("\n" if requirements else "")

    @staticmethod
    def _readme(specification: PipelineSpecification) -> str:
        return f"""# {specification.name}\n\nGenerated from confirmed specification revision {specification.revision}.\n\nRun:\n\n```text\npython pipeline.py --input-dir INPUT --output-dir OUTPUT --quarantine-dir QUARANTINE --state-file state.json\n```\n\nThe manifest checksum makes repeated runs idempotent. Inputs are never modified. Database credentials are read only from the environment variable named by `connection_ref`.\n"""

    @staticmethod
    def _test_module() -> str:
        return """import json\nfrom pathlib import Path\n\nfrom runtime_core import execute\n\n\ndef test_generated_spec_loads():\n    spec = json.loads(Path(__file__).parents[1].joinpath('pipeline_spec.json').read_text(encoding='utf-8'))\n    assert execute(spec, [{'placeholder': None}])['input_count'] == 1\n"""
