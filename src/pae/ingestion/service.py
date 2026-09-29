"""Secure, bounded readers for MVP file formats."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import re
import zipfile
from collections.abc import Iterable, Mapping, Sequence
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path, PurePosixPath
from typing import Any
from uuid import uuid4

import openpyxl  # type: ignore[import-untyped]
import pyarrow as pa  # type: ignore[import-untyped]
import pyarrow.parquet as pq  # type: ignore[import-untyped]

from pae.domain.enums import DataType
from pae.domain.models import FileDiscoveryContract
from pae.ingestion.models import (
    ColumnProfile,
    FileFormat,
    IngestionAnalysis,
    IngestionOptions,
    NormalizedFileMetadata,
    RuntimeSuggestion,
    StoredIngestion,
    WorkbookInspection,
    WorkbookSheetSummary,
)
from pae.persistence.errors import (
    ApplicationError,
    EncodingInvalid,
    FileContentInvalid,
    FileTypeUnsupported,
    IngestionLimitExceeded,
    ResourceNotFound,
    SheetSelectionRequired,
    UnsafeUpload,
    UploadTooLarge,
)

_EXTENSION_FORMAT: dict[str, FileFormat] = {
    ".csv": "csv",
    ".json": "json",
    ".jsonl": "json_lines",
    ".ndjson": "json_lines",
    ".xlsx": "excel",
    ".parquet": "parquet",
}
_ALLOWED_MIME = {
    "csv": {"text/csv", "text/plain", "application/csv", "application/vnd.ms-excel"},
    "json": {"application/json", "text/json", "text/plain"},
    "json_lines": {"application/x-ndjson", "application/jsonl", "application/json", "text/plain"},
    "excel": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    "parquet": {"application/vnd.apache.parquet", "application/octet-stream"},
}
_SAFE_NAME = re.compile(r"^[^<>:\"/\\|?*\x00-\x1f]+$")


class FileIngestionService:
    """Validate, profile, and store one temporary upload under a random name."""

    def __init__(
        self,
        storage_root: Path,
        *,
        max_file_bytes: int = 50 * 1024 * 1024,
        max_sample_rows: int = 100_000,
        max_columns: int = 1_000,
        retention_hours: int = 24,
        max_archive_entries: int = 10_000,
        max_archive_uncompressed_bytes: int = 200 * 1024 * 1024,
    ) -> None:
        self.storage_root = storage_root.resolve()
        self.max_file_bytes = max_file_bytes
        self.max_sample_rows = max_sample_rows
        self.max_columns = max_columns
        self.retention_hours = retention_hours
        self.max_archive_entries = max_archive_entries
        self.max_archive_uncompressed_bytes = max_archive_uncompressed_bytes

    def ingest(
        self,
        filename: str,
        content: bytes,
        content_type: str | None,
        options: IngestionOptions,
    ) -> StoredIngestion:
        safe_name, suffix, file_format = self._validate_envelope(filename, content, content_type)
        row_limit = min(options.sample_row_limit, self.max_sample_rows)
        rows, details = self._read(file_format, content, options, row_limit)
        columns = self._profile(rows)
        if not columns:
            raise FileContentInvalid(
                "No columns were found. Add a header and at least one structured column.",
                details={"guidance": "Provide a table/object file with named columns."},
            )
        if len(columns) > self.max_columns:
            raise IngestionLimitExceeded(
                f"The file has {len(columns)} columns; the limit is {self.max_columns}.",
                details={"limit": self.max_columns, "observed": len(columns)},
            )

        pattern = options.filename_pattern or f"*{suffix}"
        contract = FileDiscoveryContract(filename_pattern=pattern, recursive=options.recursive)
        required = tuple(column.name for column in columns if column.required)
        optional = tuple(column.name for column in columns if not column.required)
        runtime = RuntimeSuggestion(
            filename_pattern=contract.filename_pattern,
            recursive=contract.recursive,
            required_columns=required,
            optional_columns=optional,
        )
        sha256 = hashlib.sha256(content).hexdigest()
        metadata = NormalizedFileMetadata(
            original_name=safe_name,
            file_format=file_format,
            content_type=content_type or "application/octet-stream",
            size_bytes=len(content),
            sha256=sha256,
            encoding=details.get("encoding"),
            delimiter=details.get("delimiter"),
            sheet_name=details.get("sheet_name"),
            available_sheets=tuple(details.get("available_sheets", ())),
            sampled_rows=len(rows),
            total_rows=details.get("total_rows"),
            truncated=bool(details.get("truncated", False)),
            columns=columns,
        )
        self.storage_root.mkdir(parents=True, exist_ok=True)
        stored_name = f"{uuid4().hex}{suffix}"
        target = (self.storage_root / stored_name).resolve()
        if target.parent != self.storage_root:
            raise UnsafeUpload("The generated storage path escaped the sample directory.")
        with target.open("xb") as output:
            output.write(content)
            output.flush()
            os.fsync(output.fileno())
        expires_at = datetime.now(UTC) + timedelta(hours=self.retention_hours)
        return StoredIngestion(
            analysis=IngestionAnalysis(metadata=metadata, runtime=runtime),
            storage_ref=stored_name,
            expires_at=expires_at,
            sample_rows=tuple(rows),
        )

    def inspect_workbook(
        self,
        filename: str,
        content: bytes,
        content_type: str | None,
        *,
        sample_row_limit: int = 100,
    ) -> WorkbookInspection:
        """Return bounded, value-free summaries so a sheet can be selected after upload."""

        safe_name, _, file_format = self._validate_envelope(filename, content, content_type)
        if file_format != "excel":
            raise FileTypeUnsupported("Workbook inspection is available only for XLSX files.")
        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
        except Exception as exc:
            raise FileContentInvalid("The XLSX workbook could not be read.") from exc
        try:
            names = tuple(workbook.sheetnames)
            if not names:
                raise FileContentInvalid("The workbook does not contain a worksheet.")
            per_sheet_limit = min(
                sample_row_limit,
                self.max_sample_rows,
                max(1, self.max_sample_rows // len(names)),
            )
            summaries: list[WorkbookSheetSummary] = []
            for name in names:
                worksheet = workbook[name]
                total_rows = max(0, int(worksheet.max_row or 0) - 1)
                try:
                    rows, _ = self._read_excel_worksheet(workbook, name, per_sheet_limit)
                    profiles = self._profile(rows)
                    columns = tuple(profile.name for profile in profiles)
                    summaries.append(
                        WorkbookSheetSummary(
                            name=name,
                            visibility=worksheet.sheet_state,
                            status="ready" if columns else "empty",
                            sampled_rows=len(rows),
                            total_rows=total_rows,
                            column_count=len(columns),
                            columns=columns,
                        )
                    )
                except ApplicationError as exc:
                    summaries.append(
                        WorkbookSheetSummary(
                            name=name,
                            visibility=worksheet.sheet_state,
                            status="empty" if total_rows == 0 else "invalid",
                            sampled_rows=0,
                            total_rows=total_rows,
                            column_count=0,
                            warning=exc.message,
                        )
                    )
            return WorkbookInspection(
                original_name=safe_name,
                size_bytes=len(content),
                sha256=hashlib.sha256(content).hexdigest(),
                sheets=tuple(summaries),
            )
        except ApplicationError:
            raise
        except Exception as exc:
            raise FileContentInvalid("The XLSX workbook could not be inspected.") from exc
        finally:
            workbook.close()

    def store_pending(
        self, filename: str, content: bytes, content_type: str | None
    ) -> tuple[str, datetime]:
        """Store one validated workbook while the user chooses a worksheet."""

        _, suffix, file_format = self._validate_envelope(filename, content, content_type)
        if file_format != "excel":
            raise FileTypeUnsupported("Pending sheet selection is available only for XLSX files.")
        self.storage_root.mkdir(parents=True, exist_ok=True)
        stored_name = f"pending_{uuid4().hex}{suffix}"
        target = (self.storage_root / stored_name).resolve()
        if target.parent != self.storage_root:
            raise UnsafeUpload("The generated storage path escaped the sample directory.")
        with target.open("xb") as output:
            output.write(content)
            output.flush()
            os.fsync(output.fileno())
        return stored_name, datetime.now(UTC) + timedelta(hours=self.retention_hours)

    def read_stored(self, storage_ref: str) -> bytes:
        """Read a previously authorized pending upload from bounded local storage."""

        try:
            return self._resolve_ref(storage_ref).read_bytes()
        except FileNotFoundError as exc:
            raise ResourceNotFound("Pending workbook was not found or has expired.") from exc

    def discard(self, storage_ref: str) -> None:
        target = self._resolve_ref(storage_ref)
        target.unlink(missing_ok=True)

    def cleanup_expired(self, *, now: datetime | None = None) -> int:
        if not self.storage_root.exists():
            return 0
        cutoff = (now or datetime.now(UTC)) - timedelta(hours=self.retention_hours)
        removed = 0
        for candidate in self.storage_root.iterdir():
            if not candidate.is_file():
                continue
            modified = datetime.fromtimestamp(candidate.stat().st_mtime, tz=UTC)
            if modified <= cutoff:
                candidate.unlink()
                removed += 1
        return removed

    def _validate_envelope(
        self, filename: str, content: bytes, content_type: str | None
    ) -> tuple[str, str, FileFormat]:
        if not filename or Path(filename).name != filename or not _SAFE_NAME.fullmatch(filename):
            raise UnsafeUpload(
                "The filename is unsafe. Remove path segments and reserved characters."
            )
        suffix = Path(filename).suffix.lower()
        file_format = _EXTENSION_FORMAT.get(suffix)
        if file_format is None:
            raise FileTypeUnsupported(
                "Unsupported file extension. Use CSV, JSON, JSON Lines, XLSX, or Parquet.",
                details={"extension": suffix},
            )
        if not content:
            raise FileContentInvalid("The uploaded file is empty.")
        if len(content) > self.max_file_bytes:
            raise UploadTooLarge(
                f"The file exceeds the {self.max_file_bytes}-byte limit.",
                details={"limit": self.max_file_bytes, "observed": len(content)},
            )
        normalized_mime = (content_type or "application/octet-stream").split(";", 1)[0].lower()
        if (
            normalized_mime not in _ALLOWED_MIME[file_format]
            and normalized_mime != "application/octet-stream"
        ):
            raise FileTypeUnsupported(
                "The declared content type does not match the file extension.",
                details={"content_type": normalized_mime, "extension": suffix},
            )
        parquet_signature = content.startswith(b"PAR1") and content.endswith(b"PAR1")
        if parquet_signature and file_format != "parquet":
            raise FileTypeUnsupported("The file signature indicates Parquet, not its extension.")
        if file_format == "excel":
            if not content.startswith(b"PK"):
                raise FileContentInvalid("The XLSX signature is invalid.")
            self._validate_xlsx_archive(content)
        elif file_format == "parquet" and not (parquet_signature):
            raise FileContentInvalid("The Parquet signature is invalid.")
        elif content.startswith(b"PK"):
            raise UnsafeUpload("Archive uploads are not supported for this file type.")
        elif file_format == "csv" and content.lstrip().startswith((b"{", b"[")):
            raise FileTypeUnsupported("The file signature indicates JSON, not CSV.")
        return filename, suffix, file_format

    def _validate_xlsx_archive(self, content: bytes) -> None:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > self.max_archive_entries:
                    raise UnsafeUpload("The workbook contains too many archive entries.")
                total = 0
                for entry in entries:
                    path = PurePosixPath(entry.filename.replace("\\", "/"))
                    if path.is_absolute() or ".." in path.parts:
                        raise UnsafeUpload("The workbook contains an unsafe archive path.")
                    total += entry.file_size
                    if total > self.max_archive_uncompressed_bytes:
                        raise UnsafeUpload("The workbook expands beyond the safe size limit.")
                    if entry.compress_size and entry.file_size / entry.compress_size > 100:
                        raise UnsafeUpload("The workbook contains a suspicious compression ratio.")
                names = {entry.filename.replace("\\", "/") for entry in entries}
                if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                    raise FileContentInvalid("The ZIP content is not a valid XLSX workbook.")
        except zipfile.BadZipFile as exc:
            raise FileContentInvalid("The XLSX archive is corrupt.") from exc

    def _read(
        self,
        file_format: FileFormat,
        content: bytes,
        options: IngestionOptions,
        row_limit: int,
    ) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        try:
            if file_format == "csv":
                return self._read_csv(content, options, row_limit)
            if file_format == "json":
                return self._read_json(content, options, row_limit)
            if file_format == "json_lines":
                return self._read_json_lines(content, options, row_limit)
            if file_format == "excel":
                return self._read_excel(content, options, row_limit)
            return self._read_parquet(content, row_limit)
        except (
            FileContentInvalid,
            EncodingInvalid,
            SheetSelectionRequired,
            IngestionLimitExceeded,
        ):
            raise
        except Exception as exc:
            raise FileContentInvalid(
                "The file could not be read. Verify that it is not corrupt and "
                "matches its extension."
            ) from exc

    def _read_csv(
        self, content: bytes, options: IngestionOptions, row_limit: int
    ) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        text, encoding = self._decode(content, options.encoding)
        delimiter = options.delimiter
        if delimiter is None:
            try:
                delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
            except csv.Error:
                delimiter = ","
        reader = csv.DictReader(io.StringIO(text, newline=""), delimiter=delimiter)
        headers = self._validate_headers(reader.fieldnames)
        rows: list[dict[str, Any]] = []
        truncated = False
        for raw in reader:
            if None in raw:
                raise FileContentInvalid("A CSV row has more values than the header.")
            if len(rows) >= row_limit:
                truncated = True
                break
            rows.append({header: raw[header] for header in headers if raw.get(header) is not None})
        return rows, {
            "encoding": encoding,
            "delimiter": delimiter,
            "total_rows": None if truncated else len(rows),
            "truncated": truncated,
        }

    def _read_json(
        self, content: bytes, options: IngestionOptions, row_limit: int
    ) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        text, encoding = self._decode(content, options.encoding)
        try:
            parsed = json.loads(text)
        except json.JSONDecodeError as exc:
            raise FileContentInvalid(
                f"JSON is invalid near line {exc.lineno}, column {exc.colno}."
            ) from exc
        values = parsed if isinstance(parsed, list) else [parsed]
        rows = self._object_rows(values, row_limit)
        return rows, {
            "encoding": encoding,
            "total_rows": len(values),
            "truncated": len(values) > row_limit,
        }

    def _read_json_lines(
        self, content: bytes, options: IngestionOptions, row_limit: int
    ) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        text, encoding = self._decode(content, options.encoding)
        values: list[object] = []
        truncated = False
        for line_number, line in enumerate(text.splitlines(), start=1):
            if not line.strip():
                continue
            if len(values) >= row_limit:
                truncated = True
                break
            try:
                values.append(json.loads(line))
            except json.JSONDecodeError as exc:
                raise FileContentInvalid(f"JSON Lines row {line_number} is invalid.") from exc
        rows = self._object_rows(values, row_limit)
        return rows, {
            "encoding": encoding,
            "total_rows": None if truncated else len(rows),
            "truncated": truncated,
        }

    def _read_excel(
        self, content: bytes, options: IngestionOptions, row_limit: int
    ) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        workbook = openpyxl.load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        try:
            sheets = tuple(workbook.sheetnames)
            if options.sheet_name is None and len(sheets) > 1:
                raise SheetSelectionRequired(
                    "The workbook has multiple sheets. Choose one sheet to analyze.",
                    details={"available_sheets": list(sheets)},
                )
            sheet_name = options.sheet_name or sheets[0]
            if sheet_name not in sheets:
                raise SheetSelectionRequired(
                    "The selected sheet does not exist.",
                    details={"available_sheets": list(sheets)},
                )
            return self._read_excel_worksheet(workbook, sheet_name, row_limit)
        finally:
            workbook.close()

    def _read_excel_worksheet(
        self, workbook: Any, sheet_name: str, row_limit: int
    ) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        worksheet = workbook[sheet_name]
        iterator = worksheet.iter_rows(values_only=True)
        first = next(iterator, None)
        headers = self._validate_headers(
            [str(value).strip() if value is not None else "" for value in first or ()]
        )
        rows: list[dict[str, Any]] = []
        truncated = False
        for values in iterator:
            if len(rows) >= row_limit:
                truncated = True
                break
            padded = (*values, *((None,) * max(0, len(headers) - len(values))))
            rows.append(dict(zip(headers, padded[: len(headers)], strict=True)))
        total_rows = max(0, int(worksheet.max_row or 0) - 1)
        return rows, {
            "sheet_name": sheet_name,
            "available_sheets": tuple(workbook.sheetnames),
            "total_rows": total_rows,
            "truncated": truncated or total_rows > len(rows),
        }

    def _read_parquet(
        self, content: bytes, row_limit: int
    ) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        parquet = pq.ParquetFile(pa.BufferReader(content))
        self._validate_headers(parquet.schema_arrow.names)
        rows: list[dict[str, Any]] = []
        for batch in parquet.iter_batches(batch_size=min(row_limit + 1, 10_000)):
            rows.extend(batch.to_pylist())
            if len(rows) > row_limit:
                break
        total_rows = parquet.metadata.num_rows
        return rows[:row_limit], {
            "total_rows": total_rows,
            "truncated": total_rows > row_limit,
        }

    @staticmethod
    def _decode(content: bytes, requested: str | None) -> tuple[str, str]:
        candidates = (requested,) if requested else ("utf-8-sig", "utf-8", "cp874", "windows-1252")
        for encoding in candidates:
            assert encoding is not None
            try:
                return content.decode(encoding), encoding
            except (UnicodeDecodeError, LookupError):
                continue
        label = requested or "supported encodings"
        raise EncodingInvalid(
            f"The file cannot be decoded as {label}.",
            details={"guidance": "Choose the source encoding and upload again."},
        )

    def _validate_headers(self, values: Sequence[str] | None) -> tuple[str, ...]:
        if not values:
            raise FileContentInvalid("The file does not contain a header row.")
        headers = tuple(str(value).strip() for value in values)
        if any(not value for value in headers):
            raise FileContentInvalid("Column names cannot be empty.")
        if len(headers) != len(set(headers)):
            raise FileContentInvalid("Column names must be unique.")
        if len(headers) > self.max_columns:
            raise IngestionLimitExceeded(
                f"The file has {len(headers)} columns; the limit is {self.max_columns}."
            )
        return headers

    @staticmethod
    def _object_rows(values: Sequence[object], row_limit: int) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        for index, value in enumerate(values[:row_limit], start=1):
            if not isinstance(value, dict) or not all(isinstance(key, str) for key in value):
                raise FileContentInvalid(
                    f"Record {index} must be a JSON object with string field names."
                )
            rows.append(dict(value))
        return rows

    @staticmethod
    def _profile(rows: Sequence[Mapping[str, Any]]) -> tuple[ColumnProfile, ...]:
        names: list[str] = []
        for row in rows:
            for name in row:
                if name not in names:
                    names.append(name)
        profiles: list[ColumnProfile] = []
        total = len(rows)
        for name in names:
            present = sum(name in row for row in rows)
            values = [row[name] for row in rows if name in row and row[name] is not None]
            nullable = any(name in row and row[name] is None for row in rows)
            profiles.append(
                ColumnProfile(
                    name=name,
                    inferred_type=FileIngestionService._infer_type(values),
                    required=present == total,
                    nullable=nullable,
                    present_percentage=(present / total * 100) if total else 0,
                )
            )
        return tuple(profiles)

    @staticmethod
    def _infer_type(values: Iterable[Any]) -> DataType:
        observed = list(values)
        if not observed:
            return DataType.STRING
        if all(isinstance(value, bool) for value in observed):
            return DataType.BOOLEAN
        if all(isinstance(value, int) and not isinstance(value, bool) for value in observed):
            return DataType.INTEGER
        if all(
            isinstance(value, int | float | Decimal)
            and not isinstance(value, bool)
            and not (isinstance(value, float) and math.isnan(value))
            for value in observed
        ):
            return DataType.DECIMAL
        if all(isinstance(value, datetime) for value in observed):
            return DataType.DATETIME
        if all(isinstance(value, date) for value in observed):
            return DataType.DATE
        if all(isinstance(value, dict | list) for value in observed):
            return DataType.JSON
        if all(isinstance(value, str) for value in observed):
            text = [value.strip() for value in observed]
            lowered = {value.lower() for value in text}
            if lowered and lowered <= {"true", "false"}:
                return DataType.BOOLEAN
            if all(re.fullmatch(r"[+-]?(0|[1-9]\d*)", value) for value in text):
                return DataType.INTEGER
            try:
                if all(value and Decimal(value).is_finite() for value in text):
                    return DataType.DECIMAL
            except InvalidOperation:
                pass
            try:
                if all("T" in value and datetime.fromisoformat(value) for value in text):
                    return DataType.DATETIME
            except ValueError:
                pass
            try:
                if all(date.fromisoformat(value) for value in text):
                    return DataType.DATE
            except ValueError:
                pass
        return DataType.STRING

    def _resolve_ref(self, storage_ref: str) -> Path:
        if Path(storage_ref).name != storage_ref:
            raise UnsafeUpload("The storage reference is invalid.")
        target = (self.storage_root / storage_ref).resolve()
        if target.parent != self.storage_root:
            raise UnsafeUpload("The storage reference escaped the sample directory.")
        return target
